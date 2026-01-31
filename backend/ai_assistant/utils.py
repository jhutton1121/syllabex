"""Utility functions for AI assistant"""
import base64
import json
from cryptography.fernet import Fernet
from django.conf import settings
import pdfplumber


def _get_fernet():
    """Get Fernet instance using Django SECRET_KEY as the encryption key"""
    # Derive a 32-byte key from SECRET_KEY
    key = base64.urlsafe_b64encode(settings.SECRET_KEY.encode()[:32].ljust(32, b'\0'))
    return Fernet(key)


def encrypt_api_key(plaintext):
    """Encrypt an API key for storage"""
    if not plaintext:
        return ''
    f = _get_fernet()
    return f.encrypt(plaintext.encode()).decode()


def decrypt_api_key(encrypted):
    """Decrypt a stored API key"""
    if not encrypted:
        return ''
    f = _get_fernet()
    return f.decrypt(encrypted.encode()).decode()


def mask_api_key(plaintext):
    """Mask an API key for display (e.g. sk-...xxxx)"""
    if not plaintext or len(plaintext) < 8:
        return '***'
    return plaintext[:3] + '...' + plaintext[-4:]


def extract_text_from_pdf(file_obj):
    """Extract text from a PDF file, including tables as markdown"""
    text_parts = []
    with pdfplumber.open(file_obj) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)
            # Extract tables and format as markdown
            tables = page.extract_tables()
            for table in tables:
                if not table:
                    continue
                md_rows = []
                for row in table:
                    cells = [str(c).strip() if c else '' for c in row]
                    md_rows.append('| ' + ' | '.join(cells) + ' |')
                    if len(md_rows) == 1:
                        md_rows.append('| ' + ' | '.join(['---'] * len(cells)) + ' |')
                if md_rows:
                    text_parts.append('\n'.join(md_rows))

    text = '\n\n'.join(text_parts)

    # OCR fallback for scanned PDFs
    if not text.strip():
        try:
            from pdf2image import convert_from_bytes
            import pytesseract
            file_obj.seek(0)
            images = convert_from_bytes(file_obj.read())
            ocr_parts = []
            for img in images:
                ocr_text = pytesseract.image_to_string(img)
                if ocr_text.strip():
                    ocr_parts.append(ocr_text)
            text = '\n\n'.join(ocr_parts)
        except ImportError:
            pass  # OCR deps not installed — frontend warning handles visibility

    return text


def extract_text_from_docx(file_obj):
    """Extract text from a DOCX file"""
    import docx
    doc = docx.Document(file_obj)
    return '\n\n'.join(p.text for p in doc.paragraphs if p.text.strip())


def extract_text_from_file(file_obj, filename):
    """Extract text from an uploaded file based on extension"""
    lower = filename.lower()
    if lower.endswith('.pdf'):
        return extract_text_from_pdf(file_obj)
    elif lower.endswith('.docx'):
        return extract_text_from_docx(file_obj)
    else:
        raise ValueError(f"Unsupported file type: {filename}. Only PDF and DOCX are supported.")


def build_system_prompt(course, syllabus_text, assignment_context):
    """Build the system prompt for OpenAI"""
    prompt = (
        "You are an AI teaching assistant that helps instructors create assignment questions. "
        "You must return your response as valid JSON with these fields:\n"
        "1. \"message\": A brief conversational response to the instructor\n"
        "2. \"questions\": An array of question objects\n"
        "3. \"assignment_metadata\": (optional) An object with suggested assignment details. "
        "Include this when the instructor's request implies a title, description, or dates. Fields:\n"
        "   - \"title\": suggested assignment title (string)\n"
        "   - \"description\": suggested assignment description (string)\n"
        "   - \"due_date\": suggested due date as ISO 8601 string, e.g. \"2025-02-15T23:59:00\" (string or null)\n"
        "   - \"start_date\": suggested start/available date as ISO 8601 string (string or null)\n"
        "   Only include fields you can reasonably infer from the instructor's request. "
        "If the instructor doesn't mention dates, titles, or descriptions, omit assignment_metadata entirely.\n\n"
        "Each question object must have:\n"
        "- \"question_type\": one of \"multiple_choice\", \"numerical\", or \"text_response\"\n"
        "- \"text\": the question text\n"
        "- \"points\": number of points (integer)\n"
        "- \"order\": 0-based index\n\n"
        "For multiple_choice questions, also include:\n"
        "- \"choices\": array of {\"text\": string, \"is_correct\": boolean, \"order\": integer}\n"
        "  (exactly one choice must have is_correct: true)\n\n"
        "For numerical questions, also include:\n"
        "- \"correct_answer_numeric\": the correct numeric answer (float)\n"
        "- \"numeric_tolerance\": acceptable tolerance (float, e.g. 0.01)\n\n"
        "For text_response questions, no additional fields are needed.\n\n"
        "If the instructor asks you to modify, refine, or discuss questions without generating new ones, "
        "return an empty \"questions\" array and put your response in \"message\".\n\n"
        "IMPORTANT: Always return valid JSON. Do not include markdown formatting or code blocks."
        "IMPORTANT: Ensure for multiple choice \"questions\" that you are not always putting the correct answer as the first choice.\n\n"
        "IMPORTANT: If you create a multiple choice question with numeric responses, ensure they are close enough to be realistic "

    )

    if course:
        prompt += f"\n\nCourse: {course.code} - {course.name}"
        if course.description:
            prompt += f"\nCourse Description: {course.description}"

    syllabus_meta = {'syllabus_chars_total': 0, 'syllabus_chars_used': 0, 'truncated': False}
    if syllabus_text:
        syllabus_meta['syllabus_chars_total'] = len(syllabus_text)
        max_chars = 50000
        truncated = syllabus_text[:max_chars]
        if len(syllabus_text) > max_chars:
            truncated += f"\n... [syllabus truncated — {len(syllabus_text) - max_chars} characters omitted]"
            syllabus_meta['truncated'] = True
        syllabus_meta['syllabus_chars_used'] = min(len(syllabus_text), max_chars)
        prompt += f"\n\nCourse Syllabus:\n{truncated}"

    if assignment_context:
        prompt += "\n\nAssignment Context:"
        if assignment_context.get('title'):
            prompt += f"\n- Title: {assignment_context['title']}"
        if assignment_context.get('type'):
            prompt += f"\n- Type: {assignment_context['type']}"
        if assignment_context.get('points_possible'):
            prompt += f"\n- Total Points: {assignment_context['points_possible']}"

    return prompt, syllabus_meta


def build_module_system_prompt(course, syllabus_text, existing_modules, mode):
    """Build the system prompt for AI module generation"""
    prompt = (
        "You are an AI teaching assistant that helps instructors create and organize course modules. "
        "You must return your response as valid JSON with these fields:\n"
        '1. "message": A brief conversational response to the instructor\n'
        '2. "modules": An array of module objects\n\n'
        "Each module object must have:\n"
        '- "title": string (e.g. "Time Value of Money")\n'
        '- "description": string (brief overview of what the module covers)\n'
        '- "start_date": date string in YYYY-MM-DD format\n'
        '- "end_date": date string in YYYY-MM-DD format\n'
        '- "order": 0-based index (integer)\n'
        '- "zoom_link": empty string (placeholder)\n'
        '- "assignments": array of placeholder assignment objects for this module\n\n'
        "Each assignment object in the assignments array must have:\n"
        '- "title": string (e.g. "Module 1 Quiz: Time Value of Money")\n'
        '- "type": one of "quiz", "test", or "homework"\n'
        '- "due_date": date string in YYYY-MM-DD format (should fall within the module date range)\n'
        '- "points_possible": integer (reasonable default, e.g. 100 for tests, 50 for quizzes, 25 for homework)\n'
        '- "description": string (brief description of the assignment)\n\n'
        "IMPORTANT: If the syllabus or user prompt mentions assessments (quizzes, tests, exams, homework, "
        "problem sets, etc.), generate placeholder assignments with sensible titles, types, and due dates "
        "within each module's date range. These are wireframe placeholders that instructors will flesh out later.\n\n"
        "IMPORTANT: Always return valid JSON. Do not include markdown formatting or code blocks.\n"
        "IMPORTANT: Ensure module dates do not overlap unless explicitly requested.\n"
        "IMPORTANT: Account for holidays, breaks, and exam periods when the user mentions them.\n"
        "IMPORTANT: Due dates for assignments should typically be near the end of the module's date range.\n"
    )

    if mode == 'edit':
        prompt += (
            "\nWhen editing existing modules, each module object must also include:\n"
            '- "id": the existing module ID (integer) if updating an existing module, or null if creating a new one\n'
            '- "_action": one of "update", "create", or "delete"\n'
            "For deleted modules, only id and _action are required.\n"
            "Return the FULL updated list of modules (including unchanged ones with _action='update').\n\n"
        )

    if course:
        prompt += f"\n\nCourse: {course.code} - {course.name}"
        if course.description:
            prompt += f"\nCourse Description: {course.description}"
        if course.start_date:
            prompt += f"\nCourse Start Date: {course.start_date}"
        if course.end_date:
            prompt += f"\nCourse End Date: {course.end_date}"

    syllabus_meta = {'syllabus_chars_total': 0, 'syllabus_chars_used': 0, 'truncated': False}
    if syllabus_text:
        syllabus_meta['syllabus_chars_total'] = len(syllabus_text)
        max_chars = 50000
        truncated = syllabus_text[:max_chars]
        if len(syllabus_text) > max_chars:
            truncated += f"\n... [syllabus truncated — {len(syllabus_text) - max_chars} characters omitted]"
            syllabus_meta['truncated'] = True
        syllabus_meta['syllabus_chars_used'] = min(len(syllabus_text), max_chars)
        prompt += f"\n\nCourse Syllabus:\n{truncated}"

    if mode == 'edit' and existing_modules:
        prompt += "\n\nExisting Modules (current state):\n"
        prompt += json.dumps(existing_modules, indent=2, default=str)
        prompt += (
            "\n\nThe user wants to modify these existing modules. "
            "Return the full updated list with _action and id fields. "
            "Use _action='delete' for modules that should be removed. "
            "Use _action='update' for modules being modified. "
            "Use _action='create' for new modules being added."
        )

    return prompt, syllabus_meta


def call_openai(messages, ai_settings):
    """Call OpenAI API and return the parsed response"""
    from openai import OpenAI

    api_key = decrypt_api_key(ai_settings.openai_api_key_encrypted)
    if not api_key:
        raise ValueError("OpenAI API key is not configured. An admin must set it in AI Settings.")

    client = OpenAI(api_key=api_key)

    response = client.chat.completions.create(
        model=ai_settings.model_name,
        messages=messages,
        max_tokens=ai_settings.max_tokens,
        temperature=0.7,
        response_format={"type": "json_object"},
    )

    content = response.choices[0].message.content
    parsed = json.loads(content)

    # Ensure expected structure
    if 'message' not in parsed:
        parsed['message'] = ''
    if 'questions' not in parsed:
        parsed['questions'] = []

    # Assign order to questions if missing
    for i, q in enumerate(parsed['questions']):
        if 'order' not in q:
            q['order'] = i

    # Pass through assignment_metadata if present
    if 'assignment_metadata' not in parsed:
        parsed['assignment_metadata'] = None

    return parsed
